import math
import re
from io import BytesIO
from functools import lru_cache
import os
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, Group, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from backend.app.schemas.generate import (
    DownloadDocxRequest,
    DownloadPdfRequest,
)


router = APIRouter(
    prefix="/api",
    tags=["paper-download"],
)


def _add_body_text(
    document: Document,
    text: str,
) -> None:
    paragraphs = re.split(
        r"\n\s*\n",
        text.strip(),
    )

    for content in paragraphs:
        content = content.strip()

        if not content:
            continue

        paragraph = document.add_paragraph(
            content
        )

        paragraph.paragraph_format.space_after = Pt(8)


@router.post("/download/docx")
def download_docx(
    request: DownloadDocxRequest,
):
    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.size = Pt(10.5)

    # 제목
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title_paragraph.add_run(
        request.draft.title
    )

    title_run.bold = True
    title_run.font.size = Pt(18)

    # 팀 정보
    team_paragraph = document.add_paragraph()

    team_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    team_run = team_paragraph.add_run(
        "Team C (류민규, 박수암, 이혜림)"
    )

    team_run.font.size = Pt(10)

    document.add_paragraph()

    # 서론
    document.add_heading(
        "서론",
        level=1,
    )
    _add_body_text(
        document,
        request.draft.introduction,
    )

    # 본론
    document.add_heading(
        "본론",
        level=1,
    )
    _add_body_text(
        document,
        request.draft.body,
    )

    # 결론
    document.add_heading(
        "결론",
        level=1,
    )
    _add_body_text(
        document,
        request.draft.conclusion,
    )

    # 근거 자료
    if request.sources:
        document.add_heading(
            "근거 자료",
            level=1,
        )

        for index, source in enumerate(
            request.sources,
            start=1,
        ):
            source_type = (
                "논문"
                if source.type == "paper"
                else "뉴스"
            )

            paragraph = document.add_paragraph()

            paragraph.add_run(
                f"{index}. [{source_type}] "
            ).bold = True

            paragraph.add_run(
                source.title
            )

            document.add_paragraph(
                source.url
            )

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    safe_title = re.sub(
        r'[\\/:*?"<>|]',
        "",
        request.draft.title,
    ).strip()

    if not safe_title:
        safe_title = "generated_paper"

    filename = f"{safe_title}_teamC.docx"
    encoded_filename = quote(filename)

    headers = {
        "Content-Disposition": (
            "attachment; "
            'filename="generated_paper.docx"; '
            f"filename*=UTF-8''{encoded_filename}"
        )
    }

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
        headers=headers,
    )


PDF_ACCENT = colors.HexColor("#6D5EFC")
PDF_ACCENT_LIGHT = colors.HexColor("#F0EDFF")
PDF_INK = colors.HexColor("#23202F")
PDF_MUTED = colors.HexColor("#8A86A3")
PDF_BORDER = colors.HexColor("#E6E3F0")
PDF_GRID = colors.HexColor("#E9E7F1")
PDF_PALETTE = [
    colors.HexColor(value)
    for value in ("#6D5EFC", "#38BDF8", "#F472B6", "#34D399", "#FBBF24")
]

# [수정] 저장소 최상위 (backend/app/api/download.py 기준 세 단계 위)
PROJECT_ROOT = Path(__file__).resolve().parents[3]

# [수정] 순위 차트 판정용
RANK_UNITS = {"위", "순위", "등"}
RANK_TITLE_PATTERN = re.compile(r"순위|rank", re.IGNORECASE)
ROTATE_LABELS_AFTER = 6


def _font_dirs() -> list[Path]:
    """저장소 글꼴 → Windows 글꼴 → 리눅스 나눔 글꼴 순서로 찾는다."""
    return [
        PROJECT_ROOT / "assets" / "fonts",
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
        Path("/usr/share/fonts/truetype/nanum"),
    ]


@lru_cache(maxsize=1)
def _pdf_korean_fonts() -> tuple[str, str]:
    """Register regular and bold Korean fonts for polished PDF output.

    맑은고딕을 먼저 찾고, 없으면 나눔고딕을 찾는다. 배포 서버(리눅스)에서는
    저장소의 assets/fonts에 NanumGothic.ttf / NanumGothicBold.ttf를 두면 된다.
    """
    candidates = (
        ("malgun.ttf", "malgunbd.ttf"),
        ("NanumGothic.ttf", "NanumGothicBold.ttf"),
    )
    for regular_file, bold_file in candidates:
        for font_dir in _font_dirs():
            regular_path = font_dir / regular_file
            bold_path = font_dir / bold_file
            if not regular_path.is_file():
                continue
            pdfmetrics.registerFont(TTFont("PaperKorean", str(regular_path)))
            if bold_path.is_file():
                pdfmetrics.registerFont(TTFont("PaperKoreanBold", str(bold_path)))
            else:
                pdfmetrics.registerFont(TTFont("PaperKoreanBold", str(regular_path)))
            return "PaperKorean", "PaperKoreanBold"

    raise HTTPException(
        status_code=503,
        detail="PDF 다운로드에 필요한 한글 TrueType 글꼴을 찾을 수 없습니다.",
    )


def _pdf_paragraphs(text: str, style: ParagraphStyle) -> list[Paragraph]:
    paragraphs = []
    for content in re.split(r"\n+", text.strip()):
        if content.strip():
            paragraphs.append(Paragraph(escape(content.strip()), style))
    return paragraphs


def _pdf_generated_date(value: str | None) -> str:
    try:
        parsed = datetime.fromisoformat((value or "").replace("Z", "+00:00"))
    except ValueError:
        parsed = datetime.now()
    return f"{parsed.year}년 {parsed.month}월 {parsed.day}일"


def _is_rank_visual(visual) -> bool:
    """단위가 위/순위이거나 제목에 순위·rank가 있고 값이 모두 1 이상 정수면 순위 차트."""
    if visual.kind not in {"bar", "line"}:
        return False
    unit = str(getattr(visual, "unit", None) or "").strip()
    title = str(getattr(visual, "title", None) or "")
    if unit not in RANK_UNITS and not RANK_TITLE_PATTERN.search(title):
        return False
    values = [float(value) for item in visual.series for value in item.values]
    return bool(values) and all(value >= 1 and value.is_integer() for value in values)


def _rank_ticks(worst: float) -> tuple[float, list[float]]:
    """1위부터 시작하는 눈금. 최하위 막대가 사라지지 않게 끝에 한 칸 여유를 둔다."""
    raw = max(worst, 2) / 4
    magnitude = 10 ** math.floor(math.log10(raw))
    step = next(m * magnitude for m in (1, 2, 2.5, 5, 10) if m * magnitude >= raw)
    step = max(1, math.ceil(step))
    top = math.ceil(worst / step) * step
    if top <= worst:
        top += step
    ticks = [1.0] + [float(value) for value in range(step, int(top) + 1, step) if value > 1]
    return float(top), ticks


def _x_label(drawing: Drawing, x: float, y: float, text: str, font: str, rotated: bool) -> None:
    """가로축 라벨. 항목이 많으면 비스듬히 기울여 잘리거나 겹치지 않게 한다."""
    if rotated:
        label = Group(String(
            0, 0, str(text)[:18], fontName=font, fontSize=7.5,
            fillColor=PDF_MUTED, textAnchor="end",
        ))
        label.translate(x, y + 6)
        label.rotate(35)
        drawing.add(label)
        return
    drawing.add(String(
        x, y, str(text)[:16], fontName=font,
        fontSize=8, fillColor=PDF_MUTED, textAnchor="middle",
    ))


def _chart_drawing(visual, regular_font: str, bold_font: str) -> Drawing | None:
    """Render validated frontend chart data as a vector PDF graphic."""
    if visual.kind == "table":
        return None

    labels = list(visual.labels)
    series = list(visual.series)
    if not labels or not series:
        return None

    if visual.kind == "pie":
        values = list(series[0].values)
        if not values or sum(values) <= 0:
            return None
        drawing = Drawing(450, 225)
        pie = Pie()
        pie.x, pie.y, pie.width, pie.height = 118, 18, 190, 190
        pie.data = values
        pie.labels = [str(label)[:18] for label in labels]
        pie.slices.fontName = regular_font
        pie.slices.fontSize = 8
        pie.slices.strokeWidth = 0.5
        for index in range(len(values)):
            pie.slices[index].fillColor = PDF_PALETTE[index % len(PDF_PALETTE)]
        drawing.add(pie)
        return drawing

    rank = _is_rank_visual(visual)
    rotated = len(labels) > ROTATE_LABELS_AFTER
    extra = 30 if rotated else 0

    width, height = 450, 225 + extra
    drawing = Drawing(width, height)
    left, bottom, plot_width, plot_height = 48, 40 + extra, 382, 150
    values = [float(value) for item in series for value in item.values]

    if rank:
        # 순위: 위쪽이 1위, 막대는 아래 기준선에서 올라가므로 1위가 가장 길다.
        maximum, ticks = _rank_ticks(max(values))
        minimum = 1.0

        def y_position(value: float) -> float:
            return bottom + plot_height * (1 - (value - minimum) / (maximum - minimum))

        base_y = bottom
    else:
        minimum = min(0.0, min(values))
        maximum = max(0.0, max(values))
        if minimum == maximum:
            maximum = minimum + 1.0
        span = maximum - minimum
        ticks = [minimum + span * tick / 4 for tick in range(5)]

        def y_position(value: float) -> float:
            return bottom + (value - minimum) / span * plot_height

        base_y = y_position(0)

    for value in ticks:
        y = y_position(value)
        drawing.add(Line(left, y, left + plot_width, y, strokeColor=PDF_GRID, strokeWidth=0.6))
        tick_text = f"{int(value)}위" if rank else f"{value:,.1f}".replace(".0", "")
        drawing.add(String(
            left - 6, y - 3, tick_text,
            fontName=regular_font, fontSize=7.5, fillColor=PDF_MUTED,
            textAnchor="end",
        ))

    group_width = plot_width / max(1, len(labels))
    label_y = bottom - 18
    if visual.kind == "bar":
        bar_width = min(34, group_width * 0.72 / max(1, len(series)))
        for label_index, label in enumerate(labels):
            group_x = left + label_index * group_width + group_width / 2
            for series_index, item in enumerate(series):
                value = float(item.values[label_index])
                x = group_x + (series_index - (len(series) - 1) / 2) * bar_width - bar_width * 0.42
                value_y = y_position(value)
                drawing.add(Rect(
                    x, min(base_y, value_y), bar_width * 0.84,
                    max(1, abs(value_y - base_y)),
                    fillColor=PDF_PALETTE[series_index % len(PDF_PALETTE)],
                    strokeColor=None,
                    rx=3, ry=3,
                ))
                value_text = (
                    f"{int(value)}위" if rank
                    else f"{value:,.2f}".rstrip("0").rstrip(".")
                )
                drawing.add(String(
                    x + bar_width * 0.42, max(base_y, value_y) + 4,
                    value_text,
                    fontName=bold_font, fontSize=7.5, fillColor=PDF_INK,
                    textAnchor="middle",
                ))
            _x_label(drawing, group_x, label_y, label, regular_font, rotated)
    else:
        for series_index, item in enumerate(series):
            points = []
            for label_index, value in enumerate(item.values):
                x = left + group_width * (label_index + 0.5)
                y = y_position(float(value))
                points.extend([x, y])
                drawing.add(Rect(
                    x - 2.5, y - 2.5, 5, 5,
                    fillColor=PDF_PALETTE[series_index % len(PDF_PALETTE)],
                    strokeColor=None,
                ))
            drawing.add(PolyLine(
                points,
                strokeColor=PDF_PALETTE[series_index % len(PDF_PALETTE)],
                strokeWidth=2,
            ))
        for label_index, label in enumerate(labels):
            _x_label(
                drawing, left + group_width * (label_index + 0.5), label_y,
                label, regular_font, rotated,
            )

    if len(series) > 1:
        legend_x = left
        legend_y = bottom + plot_height + 17
        for index, item in enumerate(series):
            drawing.add(Rect(
                legend_x, legend_y, 8, 8,
                fillColor=PDF_PALETTE[index % len(PDF_PALETTE)], strokeColor=None,
            ))
            drawing.add(String(
                legend_x + 12, legend_y, item.name[:18], fontName=regular_font,
                fontSize=7.5, fillColor=PDF_MUTED,
            ))
            legend_x += 90
    return drawing


def _visual_flowables(request: DownloadPdfRequest, styles: dict[str, ParagraphStyle]):
    # [수정] 섹션 제목·설명을 따로 두지 않고 첫 번째 도표와 한 묶음으로 만든다.
    # 그래야 쪽 끝에 제목만 남고 도표가 다음 쪽으로 넘어가는 일이 없다.
    header = [
        Spacer(1, 8),
        Paragraph("근거 자료 시각화", styles["visual_heading"]),
        Paragraph(
            "검색된 근거 자료 원문에 있는 수치와 내용만 사용해 만들었습니다.",
            styles["visual_subtitle"],
        ),
    ]
    flowables = []
    regular_font, bold_font = _pdf_korean_fonts()

    for visual_index, visual in enumerate(request.visuals):
        source = (
            request.sources[visual.source_index]
            if 0 <= visual.source_index < len(request.sources)
            else None
        )
        if visual.kind != "table" and _is_rank_visual(visual):
            unit_text = ' <font size="7" color="#8A86A3">(순위 · 위쪽이 상위)</font>'
        elif visual.kind != "table" and visual.unit:
            unit_text = f' <font size="7" color="#8A86A3">(단위: {escape(visual.unit)})</font>'
        else:
            unit_text = ""
        card = [Paragraph(
            escape(visual.title) + unit_text,
            styles["visual_title"],
        )]
        if visual.kind == "table":
            data = [
                [Paragraph(escape(column), styles["table_head"]) for column in visual.columns]
            ]
            data.extend(
                [Paragraph(escape(str(cell)), styles["table_cell"]) for cell in row]
                for row in visual.rows
            )
            column_count = max(1, len(visual.columns))
            widths = [28 * mm] + [
                (156 * mm - 28 * mm) / max(1, column_count - 1)
            ] * max(0, column_count - 1)
            table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), PDF_ACCENT_LIGHT),
                ("TEXTCOLOR", (0, 0), (-1, 0), PDF_INK),
                ("GRID", (0, 0), (-1, -1), 0.5, PDF_BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            card.append(table)
        else:
            drawing = _chart_drawing(visual, regular_font, bold_font)
            if drawing is not None:
                card.append(drawing)

        if source:
            source_type = "논문" if source.type == "paper" else "뉴스"
            card.append(Paragraph(
                f"근거 {visual.source_index + 1} · {source_type} · {escape(source.title)}",
                styles["visual_source"],
            ))
        card.append(Spacer(1, 10))

        if visual.kind == "table":
            # 표는 쪽을 넘어가며 나뉠 수 있으므로, 제목(첫 번째면 섹션 제목 포함)만 표 첫 부분과 붙인다.
            lead = header + [card[0]] if visual_index == 0 else [card[0]]
            flowables.append(KeepTogether(lead + [card[1]]) if len(visual.rows) <= 12 else KeepTogether(lead))
            if len(visual.rows) > 12:
                flowables.extend(card[1:])
            else:
                flowables.extend(card[2:])
        else:
            lead = header if visual_index == 0 else []
            flowables.append(KeepTogether(lead + card))

    if not request.visuals:
        return header
    return flowables


def _pdf_page(canvas, document, *, title: str, font_name: str) -> None:
    canvas.saveState()
    width, height = A4
    page_number = canvas.getPageNumber()
    canvas.setStrokeColor(PDF_BORDER)
    canvas.setFillColor(PDF_MUTED)
    canvas.setFont(font_name, 7.5)
    if page_number > 1:
        canvas.drawString(18 * mm, height - 12 * mm, title[:55])
        canvas.drawRightString(width - 18 * mm, height - 12 * mm, "연예 · 문화 분야 연구 초안")
        canvas.line(18 * mm, height - 14 * mm, width - 18 * mm, height - 14 * mm)
    canvas.drawCentredString(width / 2, 10 * mm, f"- {page_number} -")
    canvas.restoreState()


@router.post("/download/pdf")
def download_pdf(request: DownloadPdfRequest):
    font_name, bold_font = _pdf_korean_fonts()
    normal = ParagraphStyle(
        "PaperBody", fontName=font_name, fontSize=10.2, leading=18,
        textColor=PDF_INK, alignment=TA_JUSTIFY, firstLineIndent=10,
        spaceAfter=9, wordWrap="CJK", splitLongWords=True,
    )
    styles = {
        "normal": normal,
        "kicker": ParagraphStyle(
            "PaperKicker", parent=normal, fontSize=8.5, leading=12,
            alignment=TA_CENTER, textColor=PDF_MUTED, firstLineIndent=0,
            spaceAfter=14,
        ),
        "title": ParagraphStyle(
            "PaperTitle", parent=normal, fontName=bold_font, fontSize=19,
            leading=29, alignment=TA_CENTER, firstLineIndent=0, spaceAfter=15,
        ),
        "team": ParagraphStyle(
            "PaperTeam", parent=normal, fontSize=9.5, alignment=TA_CENTER,
            firstLineIndent=0, textColor=colors.HexColor("#3D3A52"), spaceAfter=3,
        ),
        "date": ParagraphStyle(
            "PaperDate", parent=normal, fontSize=8.5, alignment=TA_CENTER,
            firstLineIndent=0, textColor=PDF_MUTED, spaceAfter=18,
        ),
        "heading": ParagraphStyle(
            "PaperHeading", parent=normal, fontName=bold_font, fontSize=14,
            leading=21, textColor=PDF_INK, firstLineIndent=0,
            spaceBefore=16, spaceAfter=9, keepWithNext=True,
        ),
        "visual_heading": ParagraphStyle(
            "VisualHeading", parent=normal, fontName=bold_font, fontSize=13,
            firstLineIndent=0, spaceBefore=14, spaceAfter=4, keepWithNext=True,
        ),
        "visual_subtitle": ParagraphStyle(
            "VisualSubtitle", parent=normal, fontSize=8.5, textColor=PDF_MUTED,
            firstLineIndent=0, spaceAfter=10, keepWithNext=True,
        ),
        "visual_title": ParagraphStyle(
            "VisualTitle", parent=normal, fontName=bold_font, fontSize=10.5,
            firstLineIndent=0, spaceBefore=6, spaceAfter=5,
        ),
        "visual_source": ParagraphStyle(
            "VisualSource", parent=normal, fontSize=7.5, leading=11,
            textColor=PDF_ACCENT, firstLineIndent=0, spaceBefore=4,
        ),
        "table_head": ParagraphStyle(
            "TableHead", parent=normal, fontName=bold_font, fontSize=8,
            leading=12, firstLineIndent=0, alignment=TA_LEFT,
        ),
        "table_cell": ParagraphStyle(
            "TableCell", parent=normal, fontSize=7.8, leading=12,
            firstLineIndent=0, alignment=TA_LEFT,
        ),
        "source": ParagraphStyle(
            "PaperSource", parent=normal, fontSize=8.2, leading=13,
            firstLineIndent=0, spaceAfter=6,
        ),
        "notice": ParagraphStyle(
            "PaperNotice", parent=normal, fontSize=7.5, leading=12,
            alignment=TA_CENTER, firstLineIndent=0, textColor=PDF_MUTED,
            spaceBefore=15,
        ),
    }

    story = [
        Paragraph("소셜미디어 기반 글로벌 팬덤 활동이 K-POP의 세계적 확산에 미치는 영향", styles["kicker"]),
        Paragraph(escape(request.draft.title), styles["title"]),
        Paragraph("Team C (류민규, 박수암, 이혜림)", styles["team"]),
        Paragraph(_pdf_generated_date(request.generated_at), styles["date"]),
        HRFlowable(width="100%", thickness=1.1, color=PDF_INK, spaceAfter=7),
    ]

    for number, heading, body in (
        (1, "서론", request.draft.introduction),
        (2, "본론", request.draft.body),
    ):
        story.append(Paragraph(
            f'<font color="#6D5EFC">{number}.</font>&nbsp;&nbsp;{heading}',
            styles["heading"],
        ))
        story.extend(_pdf_paragraphs(body, normal))

    if request.visuals:
        story.extend(_visual_flowables(request, styles))

    story.append(Paragraph(
        '<font color="#6D5EFC">3.</font>&nbsp;&nbsp;결론', styles["heading"]
    ))
    story.extend(_pdf_paragraphs(request.draft.conclusion, normal))

    if request.sources:
        story.append(Paragraph("참고문헌", styles["heading"]))
        for index, source in enumerate(request.sources, start=1):
            source_type = "논문" if source.type == "paper" else "뉴스"
            story.append(Paragraph(
                f'[{index}] <font color="#6D5EFC">{source_type}</font> '
                f'{escape(source.title)}<br/><font size="7" color="#8A86A3">'
                f'{escape(source.url)}</font>',
                styles["source"],
            ))

    story.append(Paragraph(
        "※ 본 문서는 Team C의 논문 초안 AI-Agent(RAG + Transformer)가 생성한 초안입니다. "
        "근거 자료를 직접 확인한 뒤 사용하세요.",
        styles["notice"],
    ))

    buffer = BytesIO()
    document = BaseDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=19 * mm, bottomMargin=18 * mm, title=request.draft.title,
        author="Team C",
    )
    page_callback = lambda canvas, doc: _pdf_page(
        canvas, doc, title=request.draft.title, font_name=font_name
    )
    frame = Frame(
        document.leftMargin,
        document.bottomMargin,
        document.width,
        document.height,
        id="paper-frame",
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    document.addPageTemplates(PageTemplate(
        id="paper-page",
        frames=[frame],
        onPageEnd=page_callback,
    ))
    document.build(story)
    buffer.seek(0)

    safe_title = re.sub(r'[\\/:*?"<>|]', "", request.draft.title).strip(" .")[:80]
    filename = f"{safe_title or 'generated_paper'}_teamC.pdf"
    headers = {
        "Content-Disposition": (
            "attachment; filename=generated_paper.pdf; "
            f"filename*=UTF-8''{quote(filename)}"
        )
    }
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)